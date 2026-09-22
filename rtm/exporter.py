"""
Excel export.

Produces a workbook with three sheets:

1. ``Summary``            -- headline coverage percentages for the run.
2. ``Traceability Matrix`` -- one row per requirement -> story -> test case
   chain, with the Execution Status cell color-coded.
3. ``Coverage Gaps``       -- orphan requirements, orphan stories, and
   never-executed test cases, each in its own labeled block.

Only ``openpyxl`` is used (no pandas dependency) to keep the tool's
footprint small and its behaviour easy to reason about.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Sequence, Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from rtm.coverage import CoverageMetrics
from rtm.gaps import GapReport
from rtm.models import TraceRow

logger = logging.getLogger("rtm.exporter")

PathLike = Union[str, Path]

HEADER_FILL = PatternFill(start_color="FF1F2937", end_color="FF1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)

STATUS_FILLS = {
    "passed": PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid"),
    "failed": PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid"),
    "blocked": PatternFill(start_color="FFFFEB9C", end_color="FFFFEB9C", fill_type="solid"),
    "not_run": PatternFill(start_color="FFD9D9D9", end_color="FFD9D9D9", fill_type="solid"),
    "no_test_case": PatternFill(start_color="FFF4CCCC", end_color="FFF4CCCC", fill_type="solid"),
}

MATRIX_HEADERS = [
    "Requirement ID",
    "Requirement Title",
    "Priority",
    "Story ID",
    "Story Title",
    "Test Case ID",
    "Test Case Title",
    "Execution Status",
]


def _style_header_row(ws: Worksheet, row: int, num_cols: int) -> None:
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _autosize_columns(ws: Worksheet, widths: Sequence[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_summary_sheet(ws: Worksheet, metrics: CoverageMetrics, gap_total: int) -> None:
    ws.title = "Summary"
    ws["A1"] = "Requirements Traceability Matrix - Coverage Summary"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:B1")

    rows = [
        ("", ""),
        ("Total requirements", metrics.total_requirements),
        ("Total user stories", metrics.total_stories),
        ("Total test cases", metrics.total_test_cases),
        ("", ""),
        ("% requirements with >=1 linked story", f"{metrics.pct_requirements_with_story:.1f}%"),
        ("% stories with >=1 linked test case", f"{metrics.pct_stories_with_test:.1f}%"),
        ("% test cases executed", f"{metrics.pct_test_cases_executed:.1f}%"),
        ("% test cases passed", f"{metrics.pct_test_cases_passed:.1f}%"),
        ("", ""),
        ("Total coverage-gap findings", gap_total),
    ]
    start_row = 3
    for offset, (label, value) in enumerate(rows):
        r = start_row + offset
        ws.cell(row=r, column=1, value=label).font = Font(bold=bool(label))
        ws.cell(row=r, column=2, value=value)

    _autosize_columns(ws, [40, 18])


def _write_matrix_sheet(ws: Worksheet, rows: Sequence[TraceRow]) -> None:
    ws.title = "Traceability Matrix"
    ws.append(MATRIX_HEADERS)
    _style_header_row(ws, 1, len(MATRIX_HEADERS))
    ws.freeze_panes = "A2"

    for row in rows:
        d = row.as_dict
        ws.append(
            [
                d["requirement_id"],
                d["requirement_title"],
                d["priority"],
                d["story_id"],
                d["story_title"],
                d["test_case_id"],
                d["test_case_title"],
                d["execution_status"],
            ]
        )
        excel_row = ws.max_row
        status = d["execution_status"]
        fill = STATUS_FILLS.get(status)
        if fill is not None:
            ws.cell(row=excel_row, column=8).fill = fill

    _autosize_columns(ws, [16, 38, 10, 12, 38, 14, 38, 16])


def _write_gaps_sheet(ws: Worksheet, gap_report: GapReport) -> None:
    ws.title = "Coverage Gaps"
    current_row = 1

    def write_block(title: str, headers: Sequence[str], data_rows: Sequence[Sequence]) -> None:
        nonlocal current_row
        ws.cell(row=current_row, column=1, value=title).font = Font(bold=True, size=12)
        current_row += 1
        if not data_rows:
            ws.cell(row=current_row, column=1, value="(none found)")
            current_row += 2
            return
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=current_row, column=col, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        current_row += 1
        for data_row in data_rows:
            for col, value in enumerate(data_row, start=1):
                ws.cell(row=current_row, column=col, value=value)
            current_row += 1
        current_row += 1  # blank spacer row

    write_block(
        f"Orphan Requirements (no linked user story) - {len(gap_report.orphan_requirements)} found",
        ["Requirement ID", "Title", "Priority"],
        [(o.requirement_id, o.title, o.priority) for o in gap_report.orphan_requirements],
    )
    write_block(
        f"Orphan User Stories (no linked test case) - {len(gap_report.orphan_stories)} found",
        ["Story ID", "Requirement ID", "Title"],
        [(o.story_id, o.requirement_id, o.title) for o in gap_report.orphan_stories],
    )
    write_block(
        f"Never-Executed Test Cases (status = not_run) - {len(gap_report.never_executed_tests)} found",
        ["Test Case ID", "Story ID", "Title"],
        [(t.test_case_id, t.story_id, t.title) for t in gap_report.never_executed_tests],
    )

    _autosize_columns(ws, [40, 40, 40])


def export_to_excel(
    output_path: PathLike,
    trace_rows: Sequence[TraceRow],
    metrics: CoverageMetrics,
    gap_report: GapReport,
) -> Path:
    """Write the Summary, Traceability Matrix, and Coverage Gaps sheets."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    summary_ws = wb.active
    _write_summary_sheet(summary_ws, metrics, gap_report.total_findings)

    matrix_ws = wb.create_sheet("Traceability Matrix")
    _write_matrix_sheet(matrix_ws, trace_rows)

    gaps_ws = wb.create_sheet("Coverage Gaps")
    _write_gaps_sheet(gaps_ws, gap_report)

    wb.save(output_path)
    logger.info("Wrote Excel workbook to %s", output_path)
    return output_path
