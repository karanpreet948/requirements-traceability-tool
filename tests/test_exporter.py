from openpyxl import load_workbook

from rtm.coverage import compute_coverage_metrics
from rtm.exporter import export_to_excel
from rtm.gaps import build_gap_report
from rtm.matrix import build_traceability


def test_export_to_excel_creates_expected_sheets(
    tmp_path, sample_requirements, sample_stories, sample_test_cases
):
    rows = build_traceability(sample_requirements, sample_stories, sample_test_cases)
    metrics = compute_coverage_metrics(sample_requirements, sample_stories, sample_test_cases)
    gap_report = build_gap_report(sample_requirements, sample_stories, sample_test_cases)

    out_path = tmp_path / "out" / "rtm.xlsx"
    result_path = export_to_excel(out_path, rows, metrics, gap_report)

    assert result_path == out_path
    assert out_path.exists()

    wb = load_workbook(out_path)
    assert wb.sheetnames == ["Summary", "Traceability Matrix", "Coverage Gaps"]

    matrix_ws = wb["Traceability Matrix"]
    # header row + 5 data rows (see conftest fixture docstring)
    assert matrix_ws.max_row == 6
    header = [c.value for c in matrix_ws[1]]
    assert header[0] == "Requirement ID"
    assert header[-1] == "Execution Status"

    gaps_ws = wb["Coverage Gaps"]
    gaps_text = "\n".join(
        str(cell.value) for row in gaps_ws.iter_rows() for cell in row if cell.value is not None
    )
    assert "R3" in gaps_text  # orphan requirement
    assert "S2" in gaps_text  # orphan story
    assert "T3" in gaps_text  # never-executed test case


def test_export_to_excel_creates_parent_directories(
    tmp_path, sample_requirements, sample_stories, sample_test_cases
):
    rows = build_traceability(sample_requirements, sample_stories, sample_test_cases)
    metrics = compute_coverage_metrics(sample_requirements, sample_stories, sample_test_cases)
    gap_report = build_gap_report(sample_requirements, sample_stories, sample_test_cases)

    nested_path = tmp_path / "a" / "b" / "c" / "rtm.xlsx"
    export_to_excel(nested_path, rows, metrics, gap_report)
    assert nested_path.exists()
