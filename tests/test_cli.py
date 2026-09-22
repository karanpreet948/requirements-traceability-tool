from pathlib import Path

import yaml
from openpyxl import load_workbook

from rtm.cli import main

REPO_ROOT = Path(__file__).resolve().parent.parent


def test_cli_build_against_committed_sample_data(tmp_path, capsys):
    output = tmp_path / "rtm.xlsx"
    exit_code = main(
        [
            "build",
            "--requirements", str(REPO_ROOT / "data" / "requirements.yaml"),
            "--stories", str(REPO_ROOT / "data" / "user_stories.yaml"),
            "--tests", str(REPO_ROOT / "data" / "test_cases.yaml"),
            "--output", str(output),
        ]
    )
    assert exit_code == 0
    assert output.exists()

    captured = capsys.readouterr()
    assert "COVERAGE SUMMARY" in captured.out
    assert "% test cases passed" in captured.out

    wb = load_workbook(output)
    assert set(wb.sheetnames) == {"Summary", "Traceability Matrix", "Coverage Gaps"}


def test_cli_build_with_missing_input_file_returns_error_not_crash(tmp_path):
    exit_code = main(
        [
            "build",
            "--requirements", str(tmp_path / "nope.yaml"),
            "--stories", str(tmp_path / "nope2.yaml"),
            "--tests", str(tmp_path / "nope3.yaml"),
            "--output", str(tmp_path / "out.xlsx"),
        ]
    )
    assert exit_code == 1


def test_sample_data_has_expected_deliberate_gaps():
    """Sanity check that the committed sample fixtures still contain the
    deliberate coverage gaps the README's sample output describes."""
    reqs = yaml.safe_load((REPO_ROOT / "data" / "requirements.yaml").read_text())["requirements"]
    stories = yaml.safe_load((REPO_ROOT / "data" / "user_stories.yaml").read_text())["user_stories"]
    tests = yaml.safe_load((REPO_ROOT / "data" / "test_cases.yaml").read_text())["test_cases"]

    req_ids_with_story = {s["requirement_id"] for s in stories}
    orphan_reqs = [r["requirement_id"] for r in reqs if r["requirement_id"] not in req_ids_with_story]
    assert len(orphan_reqs) == 4

    story_ids_with_test = {t["story_id"] for t in tests}
    orphan_stories = [s["story_id"] for s in stories if s["story_id"] not in story_ids_with_test]
    assert len(orphan_stories) == 4

    never_run = [t for t in tests if t["execution_status"] == "not_run"]
    assert len(never_run) == 8
